# -*- coding: utf-8 -*-
"""
SISTEMA DE REPORTES DE LLANTAS - GHO SATURNO
Ejecutar:  streamlit run app.py
"""
import re
import unicodedata
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ======================================================================
# CONSTANTES
# ======================================================================
NAVY = "#0F2C5C"
ACCENT = "#4C9AFF"
SHEET = "Movimientos"
DATA_DIR = Path("data")
HIST_FILE = DATA_DIR / "base_historica.pkl"

NUM_COLS = ["Km actual", "Km anterior", "Rendimiento", "# Reno", "Prof. S", "Presión",
            "Prof.E", "Presión e", "Costo", "TUERCAS", "PLOMO"]
TXT_COLS = ["Autobus", "Tipo", "Medida", "Marca", "Diseño Ori.", "Dis. Reno.", "Marca Ren.",
            "Mot. Salida", "Status", "Observacion de Salidas"]
KEY_COLS = ["Autobus", "Fecha Movimiento", "Posición", "No. Eco", "Eco. Entrada", "Km actual"]

# Valores que el usuario esperaba (para el panel de validación)
ESPERADOS = {"CPK promedio": 0.03593, "Desechos": 20, "% Desecho": 27.8, "% Cumplimiento presión": 89.3}

PAGES = ["1 · Dashboard Ejecutivo", "2 · KPI Marca Original", "3 · KPI Renovado",
         "4 · Operación y Desecho", "5 · Top Observaciones"]


# ======================================================================
# LIMPIEZA Y COLUMNAS CALCULADAS
# ======================================================================
def _norm(s) -> str:
    """Mayúsculas y sin acentos, para comparar textos."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return s.upper().strip()


def categoria_salida(obs) -> str:
    if pd.isna(obs) or str(obs).strip() == "":
        return "Sin observación"
    t = _norm(obs)
    if "PICADURA" in t:
        return "Picaduras"
    if "CORTE" in t or "PERFORA" in t:
        return "Corte / Perforación"
    if "PRESION" in t:
        return "Baja Presión"
    if "CASCO" in t:
        return "Venta de Casco"
    if "STOCK" in t:
        return "Stock - Emparejamiento"
    if "ROTA" in t:
        return "Rotación programada"
    return "Otras"


def limpiar(df: pd.DataFrame) -> pd.DataFrame:
    """Limpia la hoja Movimientos: quita filas vacías, tipifica columnas, normaliza textos."""
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df[df["Autobus"].notna() & (df["Autobus"].astype(str).str.strip() != "")].copy()
    df["Fecha Movimiento"] = pd.to_datetime(df["Fecha Movimiento"], errors="coerce")
    for c in NUM_COLS:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    for c in TXT_COLS:
        if c in df.columns:
            df[c] = df[c].astype("object").where(df[c].notna(), np.nan)
            df[c] = df[c].map(lambda x: str(x).strip() if pd.notna(x) else x)
    for c in ["Marca", "Marca Ren.", "Diseño Ori.", "Dis. Reno."]:
        if c in df.columns:
            df[c] = df[c].map(lambda x: x.upper() if isinstance(x, str) else x)
    df["Posición"] = df["Posición"].map(lambda x: str(int(x)) if isinstance(x, (int, float)) and pd.notna(x) and float(x).is_integer()
                                        else (str(x).strip() if pd.notna(x) else x))
    # Columnas que pueden venir mezcladas (números y texto) -> texto para poder guardar/exportar
    for c in ["No. Eco", "Eco. Entrada"]:
        if c in df.columns:
            df[c] = df[c].map(lambda x: str(x).strip() if pd.notna(x) else x)
    return df.reset_index(drop=True)


def enriquecer(df: pd.DataFrame) -> pd.DataFrame:
    """Agrega CPK, Mes, Vida, Categoria_Salida y Presion_Estado (idempotente)."""
    df = df.copy()
    costo = df["Costo"].fillna(0)
    rend = df["Rendimiento"]
    df["CPK"] = np.where(rend > 0, costo / rend.where(rend > 0, np.nan), 0.0)
    df["CPK"] = df["CPK"].astype(float)
    df["Mes"] = df["Fecha Movimiento"].dt.month
    df["Vida"] = np.where(df["Marca Ren."] == "ORIGINAL", "1ra Vida", "Renovada")
    df["Categoria_Salida"] = df["Observacion de Salidas"].map(categoria_salida)
    p = df["Presión"]
    df["Presion_Estado"] = np.select(
        [p.isna() | (p == 0), p.between(100, 110)], ["Sin dato", "OK"], default="Fuera de Rango")
    return df


def cargar_excel(archivo) -> pd.DataFrame:
    raw = pd.read_excel(archivo, sheet_name=SHEET)
    if "Autobus" not in [str(c).strip() for c in raw.columns]:
        raise ValueError("La hoja 'Movimientos' no tiene la columna 'Autobus'.")
    return limpiar(raw)


def combinar(base: pd.DataFrame | None, nuevo: pd.DataFrame, acumular: bool):
    """Devuelve (base_resultante, filas_agregadas, duplicados_omitidos)."""
    if not acumular or base is None or base.empty:
        return nuevo, len(nuevo), 0
    total = pd.concat([base, nuevo], ignore_index=True)
    keys = [k for k in KEY_COLS if k in total.columns]
    res = total.drop_duplicates(subset=keys, keep="last").reset_index(drop=True)
    agregadas = len(res) - len(base)
    return res, agregadas, len(nuevo) - agregadas


# ======================================================================
# AGREGACIONES (funciones puras)
# ======================================================================
def semaforo_cpk(v: float) -> str:
    return "🟢" if v < 0.08 else ("🟡" if v <= 0.15 else "🔴")


def kpis(df: pd.DataFrame) -> dict:
    n = len(df)
    des = df["Mot. Salida"] == "Desecho"
    return {
        "n": n,
        "rend": df["Rendimiento"].mean(),
        "cpk": df["CPK"].mean(),
        "cpk_pond": df["Costo"].sum() / df["Rendimiento"].sum() if df["Rendimiento"].sum() else 0,
        "n_desecho": int(des.sum()),
        "pct_desecho": des.mean() * 100 if n else 0,
        "pct_reno": (df["Status"] == "Renovada").mean() * 100 if n else 0,
        "costo": df["Costo"].sum(),
        "prof_desecho": df.loc[des, "Prof. S"].mean(),
    }


def pct_cumplimiento_presion(df: pd.DataFrame) -> tuple[float, int]:
    """% en rango 100-110 PSI sobre llantas con presión medida (>0) y cuántas quedaron sin dato."""
    medidas = df[df["Presion_Estado"] != "Sin dato"]
    if medidas.empty:
        return 0.0, len(df)
    return (medidas["Presion_Estado"] == "OK").mean() * 100, int((df["Presion_Estado"] == "Sin dato").sum())


def tabla_marca(df):
    t = df.groupby("Marca").agg(**{
        "Movimientos": ("Marca", "size"),
        "Rendimiento Prom": ("Rendimiento", "mean"),
        "Costo Prom": ("Costo", "mean"),
        "CPK Prom": ("CPK", "mean"),
        "Prof S Prom": ("Prof. S", "mean"),
    }).reset_index()
    return t.sort_values("CPK Prom").reset_index(drop=True)


def tabla_reno(df):
    t = df.groupby("Marca Ren.").agg(**{
        "Movimientos": ("Marca Ren.", "size"),
        "Rendimiento Prom": ("Rendimiento", "mean"),
        "# Reno Prom": ("# Reno", "mean"),
        "CPK Prom": ("CPK", "mean"),
        "Tasa de Falla %": ("Mot. Salida", lambda s: (s == "Desecho").mean() * 100),
    }).reset_index()
    return t.sort_values("Rendimiento Prom", ascending=False).reset_index(drop=True)


def tabla_vida(df):
    return df.groupby("Vida").agg(**{
        "Movimientos": ("Vida", "size"),
        "Rendimiento Prom": ("Rendimiento", "mean"),
        "CPK Prom": ("CPK", "mean"),
    }).reset_index()


def tabla_posicion(df):
    t = df.groupby("Posición").agg(**{"Movimientos": ("Posición", "size"),
                                      "Rendimiento Prom": ("Rendimiento", "mean"),
                                      "CPK Prom": ("CPK", "mean")}).reset_index()
    t["_o"] = t["Posición"].map(lambda x: (0, int(x)) if str(x).isdigit() else (1, 0))
    t["_s"] = t["Posición"].map(lambda x: str(x))
    return t.sort_values(["_o", "_s"]).drop(columns=["_o", "_s"]).reset_index(drop=True)


def tabla_diseno(df):
    t = df.groupby("Diseño Ori.").agg(**{"Movimientos": ("Diseño Ori.", "size"),
                                         "Rendimiento Prom": ("Rendimiento", "mean"),
                                         "CPK Prom": ("CPK", "mean")}).reset_index()
    return t.sort_values("Rendimiento Prom", ascending=False).reset_index(drop=True)


def top10_peor(df):
    cols = ["Autobus", "Posición", "Marca", "Rendimiento", "Costo", "CPK", "Mot. Salida", "Fecha Movimiento"]
    return df.nsmallest(10, "Rendimiento")[cols].reset_index(drop=True)


def alerta_desperdicio(df):
    cols = ["Autobus", "Posición", "Marca", "Prof. S", "Rendimiento", "Costo", "Fecha Movimiento", "Observacion de Salidas"]
    a = df[(df["Mot. Salida"] == "Desecho") & (df["Prof. S"] > 5)]
    return a[cols].sort_values("Prof. S", ascending=False).reset_index(drop=True)


def analisis_presion(df):
    m = df[df["Presion_Estado"] != "Sin dato"]
    bajo = m[m["Presión"] < 90]
    resto = m[m["Presión"] >= 90]
    return {
        "medidas": len(m), "sin_dato": len(df) - len(m), "n_bajo": len(bajo),
        "pct_bajo": len(bajo) / len(m) * 100 if len(m) else 0,
        "rend_bajo": bajo["Rendimiento"].mean(), "rend_resto": resto["Rendimiento"].mean(),
    }


def top_observaciones(df, n=10):
    obs = df["Observacion de Salidas"].dropna().astype(str).str.strip()
    obs = obs[obs != ""]
    if obs.empty:
        return pd.DataFrame(columns=["Observación", "Frecuencia", "% del total"]), 0
    clave = obs.map(lambda s: re.sub(r"[\s\.\;,]+$", "", _norm(s)))
    t = pd.DataFrame({"clave": clave, "Observación": obs})
    g = t.groupby("clave").agg(Observación=("Observación", "first"), Frecuencia=("Observación", "size")).reset_index(drop=True)
    g["% del total"] = g["Frecuencia"] / len(obs) * 100
    g = g.sort_values(["Frecuencia", "Observación"], ascending=[False, True]).head(n).reset_index(drop=True)
    return g, len(obs)


# ======================================================================
# EXPORTACIÓN A EXCEL
# ======================================================================
def excel_bytes(hojas: dict) -> bytes:
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as w:
        for nombre, d in hojas.items():
            d = d.copy()
            for c in d.columns:  # Excel no admite tz; asegura fechas limpias
                if pd.api.types.is_datetime64_any_dtype(d[c]):
                    d[c] = d[c].dt.tz_localize(None) if d[c].dt.tz is not None else d[c]
            d.to_excel(w, sheet_name=nombre[:31], index=False)
            ws = w.sheets[nombre[:31]]
            fill = PatternFill("solid", start_color="0F2C5C")
            for j, c in enumerate(d.columns, start=1):
                cell = ws.cell(row=1, column=j)
                cell.fill, cell.font = fill, Font(bold=True, color="FFFFFF")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                largo = max([len(str(c))] + [len(str(v)) for v in d[c].head(200).tolist()])
                ws.column_dimensions[get_column_letter(j)].width = min(max(largo + 2, 10), 45)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
    return bio.getvalue()


# ======================================================================
# PERSISTENCIA
# ======================================================================
def guardar_historico(df: pd.DataFrame):
    try:
        DATA_DIR.mkdir(exist_ok=True)
        df.to_pickle(HIST_FILE)
    except Exception:
        pass  # p. ej. disco de solo lectura: la sesión sigue funcionando


def cargar_historico():
    try:
        if HIST_FILE.exists():
            return pd.read_pickle(HIST_FILE)
    except Exception:
        pass
    return None


# ======================================================================
# UI - HELPERS
# ======================================================================
CSS = f"""
<style>
[data-testid="stMetric"] {{background:{NAVY}; border:1px solid #2B5BA8; border-radius:12px; padding:14px 16px;}}
[data-testid="stMetricLabel"] p {{color:#C9D6EE; font-size:0.85rem;}}
[data-testid="stMetricValue"] {{color:#FFFFFF;}}
h1, h2, h3 {{color:#FFFFFF;}}
section[data-testid="stSidebar"] {{border-right:1px solid {NAVY};}}
</style>
"""


def fmt_km(v):
    return "—" if pd.isna(v) else f"{v:,.0f} km"


def color_cpk(v):
    if pd.isna(v):
        return ""
    if v < 0.08:
        return "background-color:#14532d;color:#dcfce7"
    if v <= 0.15:
        return "background-color:#854d0e;color:#fef9c3"
    return "background-color:#7f1d1d;color:#fee2e2"


def tabla(df: pd.DataFrame, cpk_cols=(), key=None):
    """st.dataframe (con búsqueda/orden/filtros nativos) + semáforo en columnas CPK."""
    fmt = {}
    for c in df.columns:
        if c in ("CPK", "CPK Prom"):
            fmt[c] = "{:.4f}"
        elif c in ("Rendimiento", "Rendimiento Prom", "Km actual", "Km anterior"):
            fmt[c] = "{:,.0f}"
        elif c in ("Costo", "Costo Prom"):
            fmt[c] = "${:,.2f}"
        elif c in ("Tasa de Falla %", "% del total"):
            fmt[c] = "{:.1f}%"
        elif c in ("Prof S Prom", "# Reno Prom", "Prof. S"):
            fmt[c] = "{:.1f}"
        elif c == "Fecha Movimiento":
            fmt[c] = "{:%d/%m/%Y}"
    sty = df.style.format(fmt, na_rep="—")
    for c in cpk_cols:
        if c in df.columns:
            sty = sty.map(color_cpk, subset=[c])
    st.dataframe(sty, width="stretch", hide_index=True, key=key)


def estilo_fig(fig, height=420):
    fig.update_layout(template="plotly_dark", paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)",
                      height=height, margin=dict(l=10, r=10, t=50, b=10))
    return fig


def boton_excel(hojas: dict, nombre: str, key: str):
    st.download_button(f"⬇️ Descargar reporte ({nombre}.xlsx)", data=excel_bytes(hojas), file_name=f"{nombre}.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key=key)


# ======================================================================
# PÁGINAS
# ======================================================================
def pagina_dashboard(df):
    st.header("📊 Dashboard Ejecutivo")
    k = kpis(df)
    c = st.columns(6)
    c[0].metric("Rendimiento Promedio", fmt_km(k["rend"]))
    c[1].metric(f"CPK Promedio {semaforo_cpk(k['cpk'])}", f"{k['cpk']:.5f}",
                help=f"Promedio simple de Costo/Rendimiento por movimiento. CPK ponderado (Σ costo / Σ km): {k['cpk_pond']:.5f}")
    c[2].metric("% Desecho", f"{k['pct_desecho']:.1f}%", f"{k['n_desecho']} de {k['n']}", delta_color="off")
    c[3].metric("% Renovabilidad", f"{k['pct_reno']:.1f}%", help="Movimientos con Status = Renovada / total")
    c[4].metric("Costo Total Acumulado", f"${k['costo']:,.0f}")
    c[5].metric("Prof. Remanente al Desecho", "—" if pd.isna(k["prof_desecho"]) else f"{k['prof_desecho']:.1f} mm",
                help="Milímetros de hule que quedaban al desechar. Mientras mayor, más desperdicio.")
    st.caption("Semáforo CPK: 🟢 < 0.08 · 🟡 0.08 – 0.15 · 🔴 > 0.15")

    g1, g2 = st.columns(2)
    with g1:
        t = df.groupby("Marca")["Rendimiento"].mean().sort_values().reset_index()
        fig = px.bar(t, x="Rendimiento", y="Marca", orientation="h", text_auto=",.0f",
                     title="Rendimiento promedio por Marca Original (km)", color_discrete_sequence=[ACCENT])
        st.plotly_chart(estilo_fig(fig), width="stretch")
    with g2:
        t = df["Mot. Salida"].value_counts().reset_index()
        t.columns = ["Mot. Salida", "Movimientos"]
        fig = px.pie(t, names="Mot. Salida", values="Movimientos", hole=0.55, title="Distribución de Mot. Salida")
        fig.update_traces(textinfo="percent+value")
        st.plotly_chart(estilo_fig(fig), width="stretch")

    boton_excel({"KPIs": pd.DataFrame({"KPI": ["Movimientos", "Rendimiento Promedio", "CPK Promedio", "CPK Ponderado", "% Desecho",
                                               "% Renovabilidad", "Costo Total", "Prof. Remanente al Desecho (mm)"],
                                       "Valor": [k["n"], k["rend"], k["cpk"], k["cpk_pond"], k["pct_desecho"], k["pct_reno"],
                                                 k["costo"], k["prof_desecho"]]}),
                 "Rend_por_Marca": df.groupby("Marca")["Rendimiento"].mean().reset_index(),
                 "Mot_Salida": df["Mot. Salida"].value_counts().reset_index()},
                "DASHBOARD_EJECUTIVO", "dl_p1")

    with st.expander("✅ Validación contra tus valores esperados"):
        cump, sin_dato = pct_cumplimiento_presion(df)
        cump_todas = (df["Presion_Estado"] == "OK").mean() * 100
        real = {"CPK promedio": k["cpk"], "Desechos": k["n_desecho"], "% Desecho": k["pct_desecho"], "% Cumplimiento presión": cump}
        v = pd.DataFrame({"Indicador": list(ESPERADOS), "Esperado": list(ESPERADOS.values()),
                          "En la base cargada": [round(real[x], 5) for x in ESPERADOS]})
        v["Coincide"] = [("✅" if abs(a - b) <= max(abs(a) * 0.01, 0.05) else "⚠️") for a, b in zip(v["Esperado"], v["En la base cargada"])]
        st.dataframe(v, hide_index=True, width="stretch")
        st.caption(f"Base cargada: {k['n']} movimientos. Cumplimiento de presión calculado sobre llantas con presión medida "
                   f"(excluye {sin_dato} con 0/vacío); sobre el total sería {cump_todas:.1f}%.")


def pagina_marca(df):
    st.header("🏷️ KPI por Marca Original")
    t = tabla_marca(df)
    min_mov = st.number_input("Mínimo de movimientos para considerar una marca en la conclusión", 1, 100, 5)
    tabla(t, cpk_cols=["CPK Prom"], key="t_marca")
    elegibles = t[t["Movimientos"] >= min_mov]
    if elegibles.empty:
        st.info("Ninguna marca alcanza el mínimo de movimientos.")
    else:
        mejor, peor = elegibles.iloc[0], elegibles.iloc[-1]
        st.success(f"La mejor marca es **{mejor['Marca']}** por menor CPK ({mejor['CPK Prom']:.4f} $/km, "
                   f"{int(mejor['Movimientos'])} movimientos).")
        if len(elegibles) > 1:
            st.warning(f"La marca con mayor CPK es **{peor['Marca']}** ({peor['CPK Prom']:.4f} $/km).")
    fig = px.bar(t, x="Marca", y="CPK Prom", text_auto=".4f", title="CPK promedio por marca (menor es mejor)",
                 color="CPK Prom", color_continuous_scale=["#22c55e", "#eab308", "#ef4444"], range_color=[0, 0.25])
    fig.add_hline(y=0.08, line_dash="dot", line_color="#eab308")
    fig.add_hline(y=0.15, line_dash="dot", line_color="#ef4444")
    st.plotly_chart(estilo_fig(fig), width="stretch")
    st.caption("El CPK por movimiento es 0 cuando no hubo compra (rotaciones), por lo que marcas con muchas rotaciones "
               "salen favorecidas: revisa también el número de movimientos.")
    boton_excel({"KPI_MARCA_ORIGINAL": t}, "KPI_MARCA_ORIGINAL", "dl_p2")


def pagina_renovado(df):
    st.header("♻️ KPI Renovado")
    t = tabla_reno(df)
    tabla(t, cpk_cols=["CPK Prom"], key="t_reno")
    v = tabla_vida(df).set_index("Vida")
    st.subheader("Comparativo CPK: 1ra Vida vs Renovada")
    if {"1ra Vida", "Renovada"} <= set(v.index):
        a, b = v.loc["1ra Vida", "CPK Prom"], v.loc["Renovada", "CPK Prom"]
        c1, c2, c3 = st.columns(3)
        c1.metric(f"CPK 1ra Vida {semaforo_cpk(a)}", f"{a:.4f}", f"{int(v.loc['1ra Vida', 'Movimientos'])} mov.", delta_color="off")
        c2.metric(f"CPK Renovada {semaforo_cpk(b)}", f"{b:.4f}", f"{int(v.loc['Renovada', 'Movimientos'])} mov.", delta_color="off")
        dif = (b - a) / a * 100 if a else 0
        c3.metric("Renovada vs 1ra Vida", f"{dif:+.1f}%", "más barata por km" if dif < 0 else "más cara por km", delta_color="off")
    else:
        st.info("Se necesitan movimientos de 1ra Vida y Renovada para comparar.")
    fig = px.bar(t, x="Marca Ren.", y="Rendimiento Prom", text_auto=",.0f", title="Rendimiento promedio por Marca Renovadora (km)",
                 color_discrete_sequence=[ACCENT])
    st.plotly_chart(estilo_fig(fig), width="stretch")
    if (t["Movimientos"] < 5).any():
        chicas = ", ".join(t.loc[t["Movimientos"] < 5, "Marca Ren."])
        st.caption(f"⚠️ Muestra pequeña (< 5 movimientos): {chicas}. Interpretar con cautela.")
    boton_excel({"KPI_RENOVADO": t, "Comparativo_Vida": tabla_vida(df)}, "KPI_RENOVADO", "dl_p3")


def pagina_operacion(df):
    st.header("🛠️ Operación y Desecho")
    a, b = st.columns(2)
    with a:
        st.subheader("Rendimiento por Posición")
        tp = tabla_posicion(df)
        tabla(tp, cpk_cols=["CPK Prom"], key="t_pos")
    with b:
        st.subheader("Rendimiento por Diseño Original")
        td = tabla_diseno(df)
        tabla(td, cpk_cols=["CPK Prom"], key="t_dis")

    st.subheader("Top 10 Autobuses con peor rendimiento")
    t10 = top10_peor(df)
    tabla(t10, cpk_cols=["CPK"], key="t_top10")

    st.subheader("Análisis de presión")
    p = analisis_presion(df)
    c = st.columns(4)
    c[0].metric("% llantas < 90 PSI", f"{p['pct_bajo']:.1f}%", f"{p['n_bajo']} de {p['medidas']}", delta_color="off")
    c[1].metric("Rend. prom. < 90 PSI", fmt_km(p["rend_bajo"]))
    c[2].metric("Rend. prom. ≥ 90 PSI", fmt_km(p["rend_resto"]))
    if p["n_bajo"] and p["rend_resto"]:
        c[3].metric("Impacto en rendimiento", f"{(p['rend_bajo'] - p['rend_resto']) / p['rend_resto'] * 100:+.1f}%")
    cump, sin_dato = pct_cumplimiento_presion(df)
    st.caption(f"Cumplimiento de presión (100–110 PSI): {cump:.1f}% · {sin_dato} movimientos sin presión registrada (0/vacío) "
               f"excluidos del análisis.")

    st.subheader("🚨 Alerta de desperdicio")
    al = alerta_desperdicio(df)
    if al.empty:
        st.success("Sin llantas desechadas con más de 5 mm de profundidad.")
    else:
        st.error(f"Pérdida directa: {len(al)} llantas desechadas con más de 5mm "
                 f"(costo asociado ${al['Costo'].sum():,.2f}; promedio {al['Prof. S'].mean():.1f} mm de hule sin usar)")
        tabla(al, key="t_alerta")
    boton_excel({"Posicion": tp, "Diseno_Original": td, "Top10_Peor": t10, "Alerta_Desperdicio": al},
                "OPERACION_Y_DESECHO", "dl_p4")


def pagina_observaciones(df):
    st.header("📝 Top Observaciones de Salida")
    top, total = top_observaciones(df)
    if top.empty:
        st.info("No hay observaciones en la base.")
        return
    fig = px.bar(top.sort_values("Frecuencia"), x="Frecuencia", y="Observación", orientation="h", text="Frecuencia",
                 title=f"Top 10 observaciones más comunes (de {total} observaciones)", color_discrete_sequence=[ACCENT])
    st.plotly_chart(estilo_fig(fig, 520), width="stretch")
    tabla(top, key="t_obs")
    st.subheader("Por categoría unificada")
    cat = df["Categoria_Salida"].value_counts().reset_index()
    cat.columns = ["Categoria_Salida", "Movimientos"]
    cat["% del total"] = cat["Movimientos"] / len(df) * 100
    tabla(cat, key="t_cat")
    boton_excel({"Top_Observaciones": top, "Categoria_Salida": cat}, "TOP_OBSERVACIONES", "dl_p5")


# ======================================================================
# APP
# ======================================================================
def main():
    st.set_page_config(page_title="Reportes de Llantas · GHO Saturno", page_icon="🛞", layout="wide")
    st.markdown(CSS, unsafe_allow_html=True)
    ss = st.session_state
    if "base" not in ss:
        ss["base"] = cargar_historico()

    with st.sidebar:
        st.title("🛞 Reportes de Llantas")
        pagina = st.radio("Navegación", PAGES, key="pagina")
        st.divider()
        up = st.file_uploader("Sube Matriz_Movimientos.xlsx", type=["xlsx"])
        acumular = st.checkbox("¿Acumular histórico o reemplazar? (marcado = acumular)", value=True)
        if st.button("🔄 Actualizar Informes", width="stretch"):
            ss.pop("last_file_id", None)
            st.cache_data.clear()
            st.rerun()

        if up is not None:
            fid = f"{up.name}-{up.size}-{acumular}"
            if ss.get("last_file_id") != fid:
                try:
                    nuevo = cargar_excel(up)
                    ss["base"], agregadas, dup = combinar(ss.get("base"), nuevo, acumular)
                    guardar_historico(ss["base"])
                    ss["last_file_id"] = fid
                    modo = "acumulado" if acumular else "reemplazado"
                    ss["msg"] = ("ok", f"Base {modo}: +{agregadas} movimientos nuevos, {dup} duplicados omitidos. "
                                       f"Total: {len(ss['base'])}.")
                except Exception as e:
                    ss["msg"] = ("err", f"No se pudo leer el archivo: {e}")
        if "msg" in ss:
            (st.success if ss["msg"][0] == "ok" else st.error)(ss["msg"][1])

        base = ss.get("base")
        if base is not None and not base.empty:
            st.divider()
            excluir = st.checkbox("Excluir Rendimiento atípico", value=False,
                                  help="Oculta movimientos con Rendimiento mayor al umbral (posibles errores de Km anterior).")
            umbral = st.number_input("Umbral (km)", 100_000, 5_000_000, 500_000, step=50_000, disabled=not excluir)
            fmin, fmax = base["Fecha Movimiento"].min(), base["Fecha Movimiento"].max()
            st.caption(f"{len(base)} movimientos · {fmin:%d/%m/%Y} a {fmax:%d/%m/%Y}")
            st.download_button("⬇️ Descargar BASE_LIMPIA.xlsx", data=excel_bytes({"BASE_LIMPIA": enriquecer(base)}),
                               file_name="BASE_LIMPIA.xlsx", key="dl_base",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            if st.button("🗑️ Vaciar base guardada", width="stretch"):
                ss["base"] = None
                ss.pop("last_file_id", None)
                ss.pop("msg", None)
                HIST_FILE.unlink(missing_ok=True)
                st.rerun()

    base = ss.get("base")
    st.title("Sistema de Reportes de Llantas · GHO Saturno")
    if base is None or base.empty:
        st.info("👈 Sube **Matriz_Movimientos.xlsx** (hoja *Movimientos*) en la barra lateral para generar los informes.")
        return

    df = enriquecer(base)
    if excluir:
        n0 = len(df)
        df = df[df["Rendimiento"] <= umbral]
        st.warning(f"Se excluyeron {n0 - len(df)} movimientos con Rendimiento > {umbral:,.0f} km.")
    elif (df["Rendimiento"] > 500_000).any():
        st.caption(f"ℹ️ {(df['Rendimiento'] > 500_000).sum()} movimientos tienen Rendimiento > 500,000 km (posibles errores de captura). "
                   "Puedes excluirlos en la barra lateral.")

    {PAGES[0]: pagina_dashboard, PAGES[1]: pagina_marca, PAGES[2]: pagina_renovado,
     PAGES[3]: pagina_operacion, PAGES[4]: pagina_observaciones}[pagina](df)


if __name__ == "__main__":
    main()
